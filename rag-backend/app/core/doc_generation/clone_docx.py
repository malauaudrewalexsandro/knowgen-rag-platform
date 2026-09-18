"""DOCX clone generation (python-docx) from ordered document blocks."""
import io

from docx import Document
from docx.shared import Cm

from app.core.doc_generation.document_blocks import parse_markdown_table


def generate_docx_clone(title: str, blocks: list[dict], output_path: str) -> str:
    doc = Document()
    doc.add_heading(title, level=0)

    for b in blocks:
        if b["type"] == "heading":
            level = min(b.get("level", 1), 4)
            doc.add_heading(b["text"], level=level)
        elif b["type"] == "paragraph":
            doc.add_paragraph(b["text"])
        elif b["type"] == "table":
            columns, rows = parse_markdown_table(b["markdown"])
            if not columns:
                continue
            table = doc.add_table(rows=1, cols=len(columns))
            table.style = "Light Grid Accent 1"
            hdr_cells = table.rows[0].cells
            for i, col in enumerate(columns):
                hdr_cells[i].text = str(col)
            for row in rows:
                cells = table.add_row().cells
                for i, val in enumerate(row):
                    if i < len(cells):
                        cells[i].text = str(val)
        elif b["type"] == "image":
            try:
                doc.add_picture(io.BytesIO(b["image_bytes"]), width=Cm(14))
            except Exception:
                pass

    doc.save(output_path)
    return output_path
