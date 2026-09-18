"""
PPTX report generation (python-pptx). One slide per section: heading,
narrative textbox, then either a chart image or a data table below it.
"""
from pptx import Presentation
from pptx.util import Cm, Pt

from app.core.doc_generation.chart_builder import build_chart


def _add_table(slide, rows: list[dict], left, top, width, height):
    columns = list(rows[0].keys())
    n_rows = min(len(rows), 10) + 1  # cap visible rows so it fits a slide
    table_shape = slide.shapes.add_table(n_rows, len(columns), left, top, width, height)
    table = table_shape.table
    for i, col in enumerate(columns):
        table.cell(0, i).text = str(col)
    for r_idx, row in enumerate(rows[: n_rows - 1], start=1):
        for c_idx, col in enumerate(columns):
            table.cell(r_idx, c_idx).text = str(row.get(col, ""))


def generate_pptx(spec: dict, output_path: str) -> str:
    prs = Presentation()

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = spec["title"]

    for section in spec["sections"]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # title-only layout
        slide.shapes.title.text = section["heading"]

        top = Cm(3)
        if section.get("narrative"):
            tb = slide.shapes.add_textbox(Cm(1), top, Cm(11), Cm(2.5))
            tb.text_frame.text = section["narrative"]
            for p in tb.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(14)
            top = Cm(5.8)

        if section.get("chart_data"):
            cd = section["chart_data"]
            chart_path = build_chart(cd["type"], cd["labels"], cd["values"], title=cd.get("title"))
            slide.shapes.add_picture(chart_path, Cm(1), top, width=Cm(11))
        else:
            rows = (section.get("table_preview") or {}).get("rows") or []
            if rows:
                _add_table(slide, rows, Cm(1), top, Cm(11), Cm(8))

    prs.save(output_path)
    return output_path
