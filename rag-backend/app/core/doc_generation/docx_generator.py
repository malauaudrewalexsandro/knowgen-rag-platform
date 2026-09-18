"""DOCX report generation (python-docx)."""
from docx import Document
from docx.shared import Cm

from app.core.doc_generation.chart_builder import build_chart


def generate_docx(spec: dict, output_path: str) -> str:
    doc = Document()
    doc.add_heading(spec["title"], level=0)

    for section in spec["sections"]:
        doc.add_heading(section["heading"], level=1)
        if section.get("narrative"):
            doc.add_paragraph(section["narrative"])

        if section.get("chart_data"):
            cd = section["chart_data"]
            chart_path = build_chart(cd["type"], cd["labels"], cd["values"], title=cd.get("title"))
            doc.add_picture(chart_path, width=Cm(14))

        rows = (section.get("table_preview") or {}).get("rows") or []
        if rows:
            columns = list(rows[0].keys())
            table = doc.add_table(rows=1, cols=len(columns))
            table.style = "Light Grid Accent 1"
            hdr_cells = table.rows[0].cells
            for i, col in enumerate(columns):
                hdr_cells[i].text = str(col)
            for row in rows:
                cells = table.add_row().cells
                for i, col in enumerate(columns):
                    cells[i].text = str(row.get(col, ""))

    doc.save(output_path)
    return output_path
