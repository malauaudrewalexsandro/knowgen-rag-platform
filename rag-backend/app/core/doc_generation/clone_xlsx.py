"""
XLSX clone generation (openpyxl) from ordered document blocks. Each table
block becomes its own sheet (real, usable data); heading/paragraph text is
collected into a single "Content" sheet as rows, since prose doesn't map
cleanly to a spreadsheet — images are skipped for the same reason.
"""
from openpyxl import Workbook
from openpyxl.styles import Font

from app.core.doc_generation.document_blocks import parse_markdown_table


def generate_xlsx_clone(title: str, blocks: list[dict], output_path: str) -> str:
    wb = Workbook()
    content_ws = wb.active
    content_ws.title = "Content"
    content_ws["A1"] = title
    content_ws["A1"].font = Font(bold=True, size=14)

    row = 3
    table_count = 0
    for b in blocks:
        if b["type"] == "heading":
            content_ws.cell(row=row, column=1, value=b["text"]).font = Font(bold=True, size=12)
            row += 2
        elif b["type"] == "paragraph":
            content_ws.cell(row=row, column=1, value=b["text"])
            row += 2
        elif b["type"] == "table":
            columns, rows = parse_markdown_table(b["markdown"])
            if not columns:
                continue
            table_count += 1
            sheet_name = f"Table {table_count}"[:31]
            ws = wb.create_sheet(sheet_name)
            for c_idx, col in enumerate(columns, start=1):
                ws.cell(row=1, column=c_idx, value=col).font = Font(bold=True)
            for r_idx, data_row in enumerate(rows, start=2):
                for c_idx, val in enumerate(data_row, start=1):
                    ws.cell(row=r_idx, column=c_idx, value=val)
            content_ws.cell(row=row, column=1, value=f"[Tabel -> lihat sheet '{sheet_name}']")
            row += 2
        # images have no clean spreadsheet equivalent — skipped here.

    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    wb.save(output_path)
    return output_path
