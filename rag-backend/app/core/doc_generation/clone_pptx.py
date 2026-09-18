"""
PPTX clone generation (python-pptx) from ordered document blocks. A new
slide starts at each heading block; the content that follows it
(paragraphs/tables/images) is placed on that same slide until the next
heading, or a continuation slide if it runs out of vertical room.
"""
import io

from pptx import Presentation
from pptx.util import Cm, Pt

from app.core.doc_generation.document_blocks import parse_markdown_table

_MAX_CONTENT_TOP = Cm(16)


def _add_table(slide, columns, rows, left, top, width, height):
    n_rows = min(len(rows), 10) + 1
    table_shape = slide.shapes.add_table(n_rows, len(columns), left, top, width, height)
    table = table_shape.table
    for i, col in enumerate(columns):
        table.cell(0, i).text = str(col)
    for r_idx, row in enumerate(rows[: n_rows - 1], start=1):
        for c_idx, val in enumerate(row):
            if c_idx < len(columns):
                table.cell(r_idx, c_idx).text = str(val)


def generate_pptx_clone(title: str, blocks: list[dict], output_path: str) -> str:
    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = title

    state = {"slide": None, "top": Cm(3), "heading": title}

    def new_slide(heading_text: str):
        state["slide"] = prs.slides.add_slide(prs.slide_layouts[5])
        state["slide"].shapes.title.text = heading_text
        state["top"] = Cm(3)
        state["heading"] = heading_text

    for b in blocks:
        if b["type"] == "heading":
            new_slide(b["text"])
            continue

        if state["slide"] is None:
            new_slide(title)
        elif state["top"] > _MAX_CONTENT_TOP:
            new_slide(f"{state['heading']} (lanjutan)")

        slide = state["slide"]
        if b["type"] == "paragraph":
            tb = slide.shapes.add_textbox(Cm(1), state["top"], Cm(11), Cm(2.5))
            tb.text_frame.text = b["text"][:400]
            for p in tb.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(12)
            state["top"] += Cm(2.8)
        elif b["type"] == "table":
            columns, rows = parse_markdown_table(b["markdown"])
            if columns:
                _add_table(slide, columns, rows, Cm(1), state["top"], Cm(11), Cm(6))
                state["top"] += Cm(6.3)
        elif b["type"] == "image":
            try:
                slide.shapes.add_picture(io.BytesIO(b["image_bytes"]), Cm(1), state["top"], width=Cm(10))
                state["top"] += Cm(7)
            except Exception:
                pass

    prs.save(output_path)
    return output_path
