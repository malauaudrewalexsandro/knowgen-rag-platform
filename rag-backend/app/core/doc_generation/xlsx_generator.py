"""
XLSX report generation (openpyxl). One sheet per section. Unlike the other
three formats, the chart here is a NATIVE Excel chart object (not an
embedded PNG), so it stays interactive/editable when opened in Excel.
"""
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Font

_CHART_CLS = {"bar": BarChart, "line": LineChart, "pie": PieChart}


def _safe_sheet_name(name: str) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "_", name)[:31]
    return name or "Sheet"


def generate_xlsx(spec: dict, output_path: str) -> str:
    wb = Workbook()
    wb.remove(wb.active)

    for section in spec["sections"]:
        ws = wb.create_sheet(_safe_sheet_name(section["heading"]))
        ws["A1"] = section["heading"]
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = section.get("narrative", "")

        row = 4
        cd = section.get("chart_data")
        if cd:
            ws.cell(row=row, column=1, value="Label").font = Font(bold=True)
            ws.cell(row=row, column=2, value="Value").font = Font(bold=True)
            n = len(cd["labels"])
            for i, (label, value) in enumerate(zip(cd["labels"], cd["values"]), start=1):
                ws.cell(row=row + i, column=1, value=label)
                ws.cell(row=row + i, column=2, value=value)

            if cd["type"] in _CHART_CLS:
                chart = _CHART_CLS[cd["type"]]()
                chart.title = cd.get("title") or section["heading"]
                data_ref = Reference(ws, min_col=2, min_row=row, max_row=row + n)
                cats_ref = Reference(ws, min_col=1, min_row=row + 1, max_row=row + n)
                chart.add_data(data_ref, titles_from_data=True)
                chart.set_categories(cats_ref)
                ws.add_chart(chart, f"D{row}")
            row += n + 3

        rows = (section.get("table_preview") or {}).get("rows") or []
        if rows:
            ws.cell(row=row, column=1, value="Raw data preview").font = Font(bold=True, italic=True)
            row += 1
            columns = list(rows[0].keys())
            for c_idx, col in enumerate(columns, start=1):
                ws.cell(row=row, column=c_idx, value=col).font = Font(bold=True)
            for r_idx, r in enumerate(rows, start=row + 1):
                for c_idx, col in enumerate(columns, start=1):
                    ws.cell(row=r_idx, column=c_idx, value=r.get(col))

    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    wb.save(output_path)
    return output_path
